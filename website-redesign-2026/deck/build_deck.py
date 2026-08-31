import re, os, sys, subprocess, shutil, zipfile, math, copy
from pptx import Presentation
from pptx.util import Emu, Pt, Inches
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

SRC = 'deck_source.md'
TEMPLATE = 'template.pptx'
OUT = 'Uncommon_Website_Audit.pptx'
FONT = 'Libre Franklin'
BLUE = RGBColor(0x10, 0x34, 0xB3); YELLOW = RGBColor(0xFB, 0xAE, 0x40)
BLACK = RGBColor(0, 0, 0); GRAY = RGBColor(0x77, 0x7C, 0x86); PANEL = RGBColor(0xF2, 0xF4, 0xFA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

LEFT = Inches(0.38); WIDTH = Inches(12.56); TOP = Inches(1.12); BOTTOM = Inches(6.38)
AVAIL = (BOTTOM - TOP) / 12700 - 8  # pt, with a small safety margin above the footer
BODY_PT = 13; TBL_PT = 11; BUL_PT = 13; SUB_PT = 14

# ---------- parse ----------
def parse(md):
    sections = []; cur = None; page = None; i = 0
    lines = md.split('\n')
    def flush_page():
        nonlocal page
        if page: cur['pages'].append(page); page = None
    while i < len(lines):
        ln = lines[i]
        if ln.startswith('## '):
            flush_page()
            cur = {'title': ln[3:].strip(), 'pages': []}; sections.append(cur); i += 1; continue
        if ln.startswith('### '):
            flush_page()
            page = {'title': ln[4:].strip(), 'blocks': []}; i += 1; continue
        if ln.strip() == '' or ln.strip() == '---':
            i += 1; continue
        if page is None:
            page = {'title': cur['title'].split('. ', 1)[1] if '. ' in cur['title'] else cur['title'], 'blocks': []}
        if ln.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].startswith('|'):
                r = [c.strip() for c in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r'-+', c) for c in r): rows.append(r)
                i += 1
            page['blocks'].append({'t': 'table', 'rows': rows}); continue
        if re.match(r'^- ', ln):
            items = []
            while i < len(lines) and re.match(r'^- ', lines[i]):
                items.append(lines[i][2:].strip()); i += 1
            page['blocks'].append({'t': 'bullets', 'items': items}); continue
        if re.match(r'^\d+\. ', ln):
            items = []
            while i < len(lines) and re.match(r'^\d+\. ', lines[i]):
                m = re.match(r'^(\d+)\. (.*)', lines[i]); items.append((m.group(1), m.group(2).strip())); i += 1
            page['blocks'].append({'t': 'numbered', 'items': items}); continue
        # paragraph (single line in this source)
        page['blocks'].append({'t': 'para', 'text': ln.strip()}); i += 1
    flush_page()
    return sections

# ---------- measurement ----------
def plain(s): return re.sub(r'\*\*|`', '', s)
def lines_for(text, width_in, pt):
    cpl = max(10, int(width_in * 72 / (pt * 0.54)))
    return max(1, math.ceil(len(plain(text)) / cpl))
def para_h(text, pt=BODY_PT, width_in=12.56): return lines_for(text, width_in, pt) * pt * 1.22 + 5
def bullets_h(items, pt=BUL_PT): return sum(lines_for(t, 12.2, pt) * pt * 1.22 + 3 for t in items) + 6
def col_widths(rows):
    n = max(len(r) for r in rows)
    mx = [max(len(plain(r[c])) if c < len(r) else 0 for r in rows) for c in range(n)]
    w = [min(max(m, 6), 60) ** 0.8 for m in mx]
    tot = sum(w); widths = [12.56 * x / tot for x in w]
    # nudge columns wider when they hold long unbreakable words (URLs), then rescale to fit
    longest = [max((len(t) for r in rows if c < len(r) for t in plain(r[c]).split()), default=4) for c in range(n)]
    mins = [min(lw * TBL_PT * 0.62 / 72 + 0.2, 1.7) for lw in longest]
    fixed = [max(a, b) for a, b in zip(widths, mins)]
    tot = sum(fixed); return [12.56 * x / tot for x in fixed]
def row_h(row, widths, pt=TBL_PT, bold=False):
    l = max(lines_for(c, w - 0.16, pt + 0.6) for c, w in zip(row, widths)) if row else 1
    return l * pt * 1.28 + 6

# ---------- paginate ----------
def paginate(sections):
    pages = []
    for sec in sections:
        pages.append({'kind': 'divider', 'title': sec['title'], 'section': sec['title']})
        stitle = sec['title'].split('. ', 1)[1] if '. ' in sec['title'] else sec['title']
        cur = {'kind': 'white', 'title': stitle, 'section': sec['title'], 'blocks': []}; used = 0
        def newpage():
            nonlocal cur, used
            pages.append(cur)
            cur = {'kind': 'white', 'title': stitle + ' (continued)', 'section': sec['title'], 'blocks': []}; used = 0
        for pg in sec['pages']:
            sub = pg['title']
            # subheading must be followed by at least the first block
            first = pg['blocks'][0] if pg['blocks'] else None
            fh = 0
            if first:
                if first['t']=='para': fh = para_h(first['text'])
                elif first['t'] in ('bullets','numbered'):
                    it = first['items'][0]; fh = lines_for(it if isinstance(it,str) else it[1], 12.2, BUL_PT)*BUL_PT*1.22+3
                elif first['t']=='table':
                    w = col_widths(first['rows']); fh = row_h(first['rows'][0], w) + row_h(first['rows'][1], w)
            sh = SUB_PT*1.3 + 10
            if (used + sh + fh > AVAIL or sub.startswith('6.7')) and cur['blocks']: newpage()
            if sub != stitle:
                cur['blocks'].append({'t': 'sub', 'text': sub}); used += sh
            blocks = pg['blocks']
            for bi, b in enumerate(blocks):
                if b['t'] == 'para':
                    h = para_h(b['text'])
                    need = h
                    if re.fullmatch(r'\*\*[^*]+\*\*:?', b['text'].strip()) and bi + 1 < len(blocks):
                        nb = blocks[bi + 1]
                        if nb['t'] == 'para': need += para_h(nb['text'])
                        elif nb['t'] in ('bullets', 'numbered'):
                            it = nb['items'][0]; need += lines_for(it if isinstance(it, str) else it[1], 12.2, BUL_PT) * BUL_PT * 1.22 + 3
                        elif nb['t'] == 'table':
                            w = col_widths(nb['rows']); need += row_h(nb['rows'][0], w) + row_h(nb['rows'][1], w)
                    if used + need > AVAIL and cur['blocks']: newpage()
                    cur['blocks'].append(b); used += h
                elif b['t'] in ('bullets', 'numbered'):
                    batch = []
                    for it in b['items']:
                        txt = it if isinstance(it, str) else it[1]
                        h = lines_for(txt, 12.2, BUL_PT) * BUL_PT * 1.22 + 3
                        if used + h > AVAIL and (batch or cur['blocks']):
                            if batch: cur['blocks'].append({'t': b['t'], 'items': batch}); batch = []
                            newpage()
                        batch.append(it); used += h
                    if batch: cur['blocks'].append({'t': b['t'], 'items': batch}); used += 6
                elif b['t'] == 'table':
                    rows = b['rows']; widths = col_widths(rows); hdr = rows[0]
                    hh = row_h(hdr, widths); body = rows[1:]; batch = []
                    if used + hh + sum(row_h(r, widths) for r in body[:2]) > AVAIL and cur['blocks']: newpage()
                    used += hh
                    for r in body:
                        h = row_h(r, widths)
                        if used + h > AVAIL and batch:
                            cur['blocks'].append({'t': 'table', 'rows': [hdr] + batch, 'widths': widths}); batch = []
                            newpage(); used += hh
                        batch.append(r); used += h
                    cur['blocks'].append({'t': 'table', 'rows': [hdr] + batch, 'widths': widths}); used += 12
        pages.append(cur)
    return pages

# ---------- package ----------
def build_package(n_white, n_blue):
    if os.path.exists('un'): shutil.rmtree('un')
    zipfile.ZipFile(TEMPLATE).extractall('un')
    add = '/mnt/skills/public/pptx/scripts/add_slide.py'
    whites = []; blues = []
    for _ in range(n_white):
        out = subprocess.check_output(['python3', add, 'un/', 'slide2.xml'], text=True)
        whites.append(re.search(r'slides/(slide\d+)\.xml', out).group(1))
    for _ in range(n_blue):
        out = subprocess.check_output(['python3', add, 'un/', 'slideLayout8.xml'], text=True)
        blues.append(re.search(r'slides/(slide\d+)\.xml', out).group(1))
    # top lockup (logo + rule + Change History) as used on the cover; lives in layout 4's media
    l4 = open('un/ppt/slideLayouts/_rels/slideLayout4.xml.rels').read()
    lock = re.search(r'Target="\.\./media/(image\d+\.png)"', l4).group(1)
    shutil.copy(f'un/ppt/media/{lock}', 'lockup.png')
    return whites, blues

def set_order(order_names):
    pres = 'un/ppt/presentation.xml'; rels = 'un/ppt/_rels/presentation.xml.rels'
    rx = open(rels).read()
    rid = {}
    for m in re.finditer(r'<Relationship Id="(rId\d+)"[^>]*Target="slides/(slide\d+)\.xml"', rx): rid[m.group(2)] = m.group(1)
    for m in re.finditer(r'<Relationship [^>]*Target="slides/(slide\d+)\.xml"[^>]*Id="(rId\d+)"', rx): rid[m.group(1)] = m.group(2)
    px = open(pres).read()
    ids = ''.join(f'<p:sldId id="{300+i}" r:id="{rid[n]}"/>' for i, n in enumerate(order_names))
    px = re.sub(r'<p:sldIdLst>.*?</p:sldIdLst>', f'<p:sldIdLst>{ids}</p:sldIdLst>', px, flags=re.S)
    open(pres, 'w').write(px)

# ---------- fill ----------
def strip_shapes(slide, keep_pred):
    for sh in list(slide.shapes):
        if not keep_pred(sh): sh._element.getparent().remove(sh._element)

def is_title(sh): return sh.top == 274638 and sh.has_text_frame
def is_rule(sh): return sh.shape_type == 9  # line on blue page

def set_title(slide, text):
    for sh in slide.shapes:
        if is_title(sh):
            p = sh.text_frame.paragraphs[0]; r = p.runs[0]; r.text = text
            for extra in p.runs[1:]: extra._r.getparent().remove(extra._r)
            for extra in sh.text_frame.paragraphs[1:]: extra._p.getparent().remove(extra._p)
            return

def style_run(r, pt, bold=False, color=BLACK, italic=False):
    r.font.name = FONT; r.font.size = Pt(pt); r.font.bold = bold; r.font.italic = italic; r.font.color.rgb = color
    rPr = r._r.get_or_add_rPr()
    for tag in ('a:ea', 'a:cs'):
        el = rPr.find(qn(tag))
        if el is None: el = etree.SubElement(rPr, qn(tag))
        el.set('typeface', FONT)

def add_runs(p, text, pt, color=BLACK, base_bold=False):
    parts = re.split(r'(\*\*[^*]+\*\*)', text)
    for part in parts:
        if not part: continue
        b = part.startswith('**')
        t = part.strip('*').replace('`', '')
        r = p.add_run(); r.text = t; style_run(r, pt, bold=(b or base_bold), color=color)

def add_textbox(slide, x, y, w, h):
    tb = slide.shapes.add_textbox(x, y, w, h); tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = 0; tf.margin_top = tf.margin_bottom = 0
    return tf

def set_bullet(p, char='•', indent=Emu(228600)):
    pPr = p._p.get_or_add_pPr(); pPr.set('marL', str(int(indent))); pPr.set('indent', str(-int(indent)))
    bu = etree.SubElement(pPr, qn('a:buClr')); c = etree.SubElement(bu, qn('a:srgbClr')); c.set('val', '1034B3')
    bf = etree.SubElement(pPr, qn('a:buFont')); bf.set('typeface', 'Arial')
    bc = etree.SubElement(pPr, qn('a:buChar')); bc.set('char', char)

def fill_white(slide, page):
    strip_shapes(slide, is_title); set_title(slide, page['title'])
    y = TOP
    for b in page['blocks']:
        if b['t'] == 'sub':
            h = Pt(SUB_PT*1.3 + 10); tf = add_textbox(slide, LEFT, y + Pt(4), WIDTH, h)
            p = tf.paragraphs[0]; r = p.add_run(); r.text = b['text']; style_run(r, SUB_PT, bold=True, color=BLUE)
            y += h; continue
        if b['t'] == 'para':
            h = Pt(para_h(b['text'])); tf = add_textbox(slide, LEFT, y, WIDTH, h)
            p = tf.paragraphs[0]; add_runs(p, b['text'], BODY_PT); p.space_after = Pt(4)
            y += h
        elif b['t'] in ('bullets', 'numbered'):
            h = Pt(bullets_h([it if isinstance(it, str) else it[1] for it in b['items']]))
            tf = add_textbox(slide, LEFT, y, WIDTH, h); first = True
            for it in b['items']:
                p = tf.paragraphs[0] if first else tf.add_paragraph(); first = False; p.space_after = Pt(3)
                if b['t'] == 'bullets':
                    set_bullet(p); add_runs(p, it, BUL_PT)
                else:
                    set_bullet(p, char='', indent=Emu(320000)); num, txt = it
                    r = p.add_run(); r.text = f'{num}.  '; style_run(r, BUL_PT, bold=True, color=BLUE)
                    add_runs(p, txt, BUL_PT)
            y += h
        elif b['t'] == 'table':
            rows = b['rows']; widths = b['widths']; n = len(widths)
            hs = [row_h(r, widths) for r in rows]; H = Pt(sum(hs))
            gf = slide.shapes.add_table(len(rows), n, LEFT, y, WIDTH, H); tbl = gf.table
            tblPr = tbl._tbl.tblPr; tblPr.set('firstRow', '0'); tblPr.set('bandRow', '0')
            for sid in tblPr.findall(qn('a:tableStyleId')): tblPr.remove(sid)
            for ci, w in enumerate(widths): tbl.columns[ci].width = Inches(w)
            for ri, row in enumerate(rows):
                tbl.rows[ri].height = Pt(hs[ri])
                for ci in range(n):
                    cell = tbl.cell(ri, ci); txt = row[ci] if ci < len(row) else ''
                    cell.margin_left = cell.margin_right = Emu(50000); cell.margin_top = cell.margin_bottom = Emu(25000)
                    cell.vertical_anchor = MSO_ANCHOR.TOP
                    tf = cell.text_frame; p = tf.paragraphs[0]
                    add_runs(p, txt, TBL_PT, base_bold=(ri == 0))
                    cell.fill.solid(); cell.fill.fore_color.rgb = PANEL if ri == 0 else WHITE
                    tcPr = cell._tc.get_or_add_tcPr()
                    for side in ('a:lnL', 'a:lnR', 'a:lnT', 'a:lnB'):
                        ln = etree.SubElement(tcPr, qn(side)); ln.set('w', '6350')
                        sf = etree.SubElement(ln, qn('a:solidFill')); c = etree.SubElement(sf, qn('a:srgbClr')); c.set('val', 'C9CED8')
            y += H + Pt(12)

def fill_divider(slide, page, idx, total):
    strip_shapes(slide, lambda s: False)
    t = page['title']; num, name = (t.split('. ', 1) if '. ' in t else ('', t))
    # same top lockup as the cover, same position
    slide.shapes.add_picture('lockup.png', Emu(304800), Emu(325957), Emu(11558021), Emu(949044))
    # same title box as the cover: bottom-anchored, 53pt bold white, 90% line spacing
    tb = slide.shapes.add_textbox(Emu(304800), Emu(2043611), Emu(10261500), Emu(2149200)); tf = tb.text_frame
    tf.word_wrap = True; tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.BOTTOM
    p = tf.paragraphs[0]; p.line_spacing = 0.9
    r = p.add_run(); r.text = name; style_run(r, 53, bold=True, color=WHITE)

def fill_cover(slide):
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            p = sh.text_frame.paragraphs[0]; p.runs[0].text = 'Website Redesign: Technical and Content Audit'
            for extra in p.runs[1:]: extra._r.getparent().remove(extra._r)
            for extra in sh.text_frame.paragraphs[1:]: extra._p.getparent().remove(extra._p)

def fill_toc(slide, entries):
    strip_shapes(slide, is_title); set_title(slide, 'Contents')
    half = math.ceil(len(entries) / 2)
    for col, chunk in enumerate((entries[:half], entries[half:])):
        x = LEFT + col * Inches(6.45)
        tf = add_textbox(slide, x, TOP, Inches(6.3), BOTTOM - TOP); first = True
        for name, pg in chunk:
            p = tf.paragraphs[0] if first else tf.add_paragraph(); first = False; p.space_after = Pt(7)
            num, nm = (name.split('. ', 1) if '. ' in name else ('', name))
            r = p.add_run(); r.text = f'{num}.  ' if num else ''; style_run(r, 13, bold=True, color=BLUE)
            r = p.add_run(); r.text = nm; style_run(r, 13, bold=True)
            r = p.add_run(); r.text = f'   p. {pg}'; style_run(r, 11, color=GRAY)

def main():
    sections = parse(open(SRC).read())
    pages = paginate(sections)
    n_white = sum(1 for p in pages if p['kind'] == 'white') + 1  # + TOC
    n_blue = sum(1 for p in pages if p['kind'] == 'divider')
    whites, blues = build_package(n_white, n_blue)
    order = ['slide1']; wi = 0; bi = 0; toc_name = whites[wi]; order.append(toc_name); wi += 1
    slide_of = []
    for p in pages:
        if p['kind'] == 'divider': p['slide'] = blues[bi]; bi += 1
        else: p['slide'] = whites[wi]; wi += 1
        order.append(p['slide'])
    set_order(order)
    subprocess.check_call(['python3', '/mnt/skills/public/pptx/scripts/clean.py', 'un/'])
    if os.path.exists('stage.pptx'): os.remove('stage.pptx')
    subprocess.check_call('cd un && zip -Xrq ../stage.pptx .', shell=True)
    prs = Presentation('stage.pptx')
    slides = list(prs.slides)
    fill_cover(slides[0])
    entries = []
    for i, p in enumerate(pages):
        if p['kind'] == 'divider': entries.append((p['title'], i + 3))
    fill_toc(slides[1], entries)
    for i, p in enumerate(pages):
        s = slides[i + 2]
        if p['kind'] == 'divider': fill_divider(s, p, 0, 0)
        else: fill_white(s, p)
    prs.save(OUT)
    print('pages', len(pages), 'slides', len(slides))
    # pagination report
    conts = [p['title'] for p in pages if p['kind'] == 'white' and p['title'].endswith('(continued)')]
    print('continued pages:', len(conts))

if __name__ == '__main__': main()
